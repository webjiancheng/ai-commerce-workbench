from __future__ import annotations

import json
from functools import lru_cache
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


TEMPLATE_SHEET_NAME = "模版"
TEMPLATE_COMMON_ROW = 1
TEMPLATE_COMMON_VALUES_ROW = 2
TEMPLATE_GROUP_ROW = 3
TEMPLATE_HEADER_ROW = 4
TEMPLATE_HINT_ROW = 5
TEMPLATE_DATA_START_ROW = 6

# 旧字段兼容映射。读取时兼容，写入时统一输出为右侧字段。
LEGACY_FIELD_ALIASES: dict[str, str] = {
    "product_title_cn": "商品名称",
    "product_title_en": "英文名称",
    "category_path": "类目",
    "申报价CNY": "申报价格-美国站",
    "建议售价CNY": "制造商建议零售价(USD)",
    "库存": "发货仓1库存",
    "商品轮播图1-英语": "商品轮播图1",
    "商品轮播图2-英语": "商品轮播图2",
    "商品轮播图3-英语": "商品轮播图3",
    "商品轮播图4-英语": "商品轮播图4",
    "商品轮播图5-英语": "商品轮播图5",
    "商品轮播图6-英语": "商品轮播图6",
    "商品轮播图7-英语": "商品轮播图7",
    "商品轮播图8-英语": "商品轮播图8",
    "商品轮播图9-英语": "商品轮播图9",
    "商品轮播图10-英语": "商品轮播图10",
    "SPU主图视频": "主图视频",
    "SPU详情视频": "详情视频",
    "产品轮播图": "商品轮播图1",
    "产品素材图": "SKU预览图-英语",
}

# 抽屉仍需要的辅助字段；这些不是模板字段，但要保留。
AUXILIARY_FIELD_KEYS: set[str] = {
    "商品属性明细",
    "SKU信息明细",
    "敏感属性明细",
}


def get_builtin_temu_upload_template_path() -> str:
    return str(Path(__file__).resolve().parents[4] / "商品上传模版.xlsx")


def normalize_listing_field_key(key: str | None) -> str:
    raw = str(key or "").strip()
    if not raw:
        return ""
    return LEGACY_FIELD_ALIASES.get(raw, raw)


def get_legacy_aliases_for_field(field_name: str) -> list[str]:
    normalized = normalize_listing_field_key(field_name)
    aliases = [key for key, value in LEGACY_FIELD_ALIASES.items() if value == normalized]
    return sorted(set(aliases))


def normalize_listing_fields(
    fields: dict[str, Any] | None,
    *,
    keep_unknown: bool = True,
    include_auxiliary: bool = True,
) -> dict[str, Any]:
    normalized: dict[str, Any] = {}
    if not isinstance(fields, dict):
        return normalized

    meta = get_temu_upload_template_meta()
    allowed = set(meta["all_field_names"])
    if include_auxiliary:
        allowed |= AUXILIARY_FIELD_KEYS

    for key, value in fields.items():
        normalized_key = normalize_listing_field_key(str(key))
        if not normalized_key:
            continue
        if not keep_unknown and normalized_key not in allowed:
            continue
        normalized[normalized_key] = value
    return normalized


def read_listing_field(fields: dict[str, Any] | None, field_name: str, default: Any = None) -> Any:
    normalized = normalize_listing_field_key(field_name)
    if not isinstance(fields, dict):
        return default
    if normalized in fields:
        return fields[normalized]
    for alias in get_legacy_aliases_for_field(normalized):
        if alias in fields:
            return fields[alias]
    return default


@lru_cache(maxsize=1)
def get_temu_upload_template_meta() -> dict[str, Any]:
    return parse_template_meta(template_file_path=get_builtin_temu_upload_template_path())


def parse_template_meta(*, template_file_path: str) -> dict[str, Any]:
    wb = load_workbook(template_file_path, data_only=True)
    ws = wb[TEMPLATE_SHEET_NAME] if TEMPLATE_SHEET_NAME in wb.sheetnames else wb[wb.sheetnames[0]]

    common_fields: list[dict[str, Any]] = []
    detail_fields: list[dict[str, Any]] = []
    required_fields: list[str] = []
    conditional_required_fields: list[str] = []
    required_hints: list[str] = []
    enum_options_map: dict[str, list[str]] = {}

    current_group = ""
    header_key_map: dict[str, str] = {}
    key_value_map: dict[str, list[str]] = {}

    if "HeaderKeyMap" in wb.sheetnames:
        header_key_ws = wb["HeaderKeyMap"]
        for col_idx in range(1, ws.max_column + 1):
            raw_value = header_key_ws.cell(row=1, column=col_idx).value
            key = str(raw_value or "").strip()
            if key:
                header_key_map[key] = str(ws.cell(row=TEMPLATE_HEADER_ROW, column=col_idx).value or "").strip()

    if "KeyValueMap" in wb.sheetnames:
        key_value_ws = wb["KeyValueMap"]
        for col_idx in range(1, ws.max_column + 1):
            raw_header_key = (
                header_key_ws.cell(row=1, column=col_idx).value
                if "HeaderKeyMap" in wb.sheetnames
                else None
            )
            header_key = str(raw_header_key or "").strip()
            if not header_key:
                continue
            raw_value = key_value_ws.cell(row=1, column=col_idx).value
            options = _parse_key_value_map_cell(raw_value)
            if options:
                key_value_map[header_key] = options

    conditional_rules = _parse_property_relate_require(wb, header_key_map)

    for col_idx in range(1, ws.max_column + 1):
        common_name = str(ws.cell(row=TEMPLATE_COMMON_ROW, column=col_idx).value or "").strip()
        common_default = str(ws.cell(row=TEMPLATE_COMMON_VALUES_ROW, column=col_idx).value or "").strip()
        if common_name:
            common_fields.append(
                {
                    "field_name": common_name,
                    "column_index": col_idx,
                    "column_letter": get_column_letter(col_idx),
                    "default_value": common_default,
                }
            )

        group_name = str(ws.cell(row=TEMPLATE_GROUP_ROW, column=col_idx).value or "").strip()
        if group_name:
            current_group = group_name

        header_name = str(ws.cell(row=TEMPLATE_HEADER_ROW, column=col_idx).value or "").strip()
        if not header_name:
            continue

        hint = str(ws.cell(row=TEMPLATE_HINT_ROW, column=col_idx).value or "").strip()
        required, sku_required, conditional_required = _classify_hint_requirement(hint)

        if required or sku_required:
            required_fields.append(header_name)
        if conditional_required:
            conditional_required_fields.append(header_name)
        if hint:
            required_hints.append(f"{header_name}: {hint}")

        header_key = next((key for key, name in header_key_map.items() if name == header_name), "")
        options = key_value_map.get(header_key, [])
        if options:
            enum_options_map[header_name] = options

        detail_fields.append(
            {
                "field_name": header_name,
                "column_index": col_idx,
                "column_letter": get_column_letter(col_idx),
                "group": current_group,
                "required": required or sku_required,
                "conditional_required": conditional_required,
                "hint": hint,
                "enum_options": options,
            }
        )

    return {
        "adapter_key": "temu_half_managed_jewelry_upload",
        "template_file_path": template_file_path,
        "sheet_name": ws.title,
        "common_row": TEMPLATE_COMMON_ROW,
        "common_values_row": TEMPLATE_COMMON_VALUES_ROW,
        "group_row": TEMPLATE_GROUP_ROW,
        "header_row": TEMPLATE_HEADER_ROW,
        "hint_row": TEMPLATE_HINT_ROW,
        "data_start_row": TEMPLATE_DATA_START_ROW,
        "common_fields": common_fields,
        "detail_fields": detail_fields,
        "required_fields": sorted(set(required_fields)),
        "conditional_required_fields": sorted(set(conditional_required_fields)),
        "required_hints": required_hints,
        "enum_options_map": enum_options_map,
        "conditional_rules": conditional_rules,
        "groups": _build_groups(detail_fields),
        "detail_headers": [item["field_name"] for item in detail_fields],
        "common_field_names": [item["field_name"] for item in common_fields],
        "default_values": {item["field_name"]: item.get("default_value", "") for item in common_fields},
        "all_field_names": [item["field_name"] for item in common_fields] + [item["field_name"] for item in detail_fields],
        "auxiliary_field_keys": sorted(AUXILIARY_FIELD_KEYS),
        "legacy_field_aliases": dict(LEGACY_FIELD_ALIASES),
    }


def _build_groups(detail_fields: list[dict[str, Any]]) -> list[dict[str, Any]]:
    groups: list[dict[str, Any]] = []
    current_name = ""
    current_fields: list[str] = []
    for item in detail_fields:
        name = str(item.get("group") or "").strip() or "未分组"
        field_name = str(item["field_name"])
        if name != current_name:
            if current_fields:
                groups.append({"name": current_name, "fields": current_fields})
            current_name = name
            current_fields = [field_name]
        else:
            current_fields.append(field_name)
    if current_fields:
        groups.append({"name": current_name, "fields": current_fields})
    return groups


def _classify_hint_requirement(hint: str) -> tuple[bool, bool, bool]:
    text = str(hint or "").strip()
    if not text:
        return False, False, False
    if text.startswith("必填"):
        return True, False, False
    if "条件必填" in text:
        return False, False, True
    if "SKU层必填" in text:
        return False, True, False
    if "非必填" in text or "条件选填" in text or "选填" in text:
        return False, False, False
    if any(token in text for token in ["若", "如果", "如"]) and any(token in text for token in ["必填", "必须填写"]):
        return False, False, True
    if "必填" in text:
        return True, False, False
    return False, False, False


def _parse_key_value_map_cell(raw_value: Any) -> list[str]:
    text = str(raw_value or "").strip()
    if not text:
        return []
    try:
        parsed = json.loads(text)
    except Exception:
        return []
    if isinstance(parsed, dict):
        return [str(key).strip() for key in parsed.keys() if str(key).strip()]
    if isinstance(parsed, list):
        return [str(item).strip() for item in parsed if str(item).strip()]
    return []


def _parse_property_relate_require(workbook: Any, header_key_map: dict[str, str]) -> list[dict[str, Any]]:
    if "PropertyRelateRequire" not in workbook.sheetnames:
        return []

    relate_ws = workbook["PropertyRelateRequire"]
    rules: list[dict[str, Any]] = []

    for row_idx in range(1, relate_ws.max_row + 1):
        trigger_value = str(relate_ws.cell(row=row_idx, column=1).value or "").strip()
        if not trigger_value or "_" not in trigger_value:
            continue
        option_label, header_key = trigger_value.rsplit("_", 1)
        trigger_field_name = header_key_map.get(header_key)
        if not trigger_field_name:
            continue

        required_headers: list[str] = []
        for col_idx in range(1, min(relate_ws.max_column, len(header_key_map)) + 1):
            cell_value = str(relate_ws.cell(row=row_idx, column=col_idx).value or "").strip().lower()
            if cell_value != "require":
                continue
            raw_header_key = None
            if "HeaderKeyMap" in workbook.sheetnames:
                raw_header_key = workbook["HeaderKeyMap"].cell(row=1, column=col_idx).value
            header_name = header_key_map.get(str(raw_header_key or "").strip())
            if header_name:
                required_headers.append(header_name)

        if required_headers:
            rules.append(
                {
                    "if_field": trigger_field_name,
                    "if_value": option_label,
                    "required_headers": sorted(set(required_headers)),
                }
            )
    return rules
