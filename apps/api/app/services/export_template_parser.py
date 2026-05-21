from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from slugify import slugify


@dataclass(frozen=True)
class TemplateField:
    field_key: str
    field_name: str
    column_index: int
    required: bool


def parse_template_fields(*, file_path: str, header_row_index: int = 1, sheet_name: str = "Sheet1") -> list[TemplateField]:
    wb = load_workbook(file_path)
    ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb[wb.sheetnames[0]]

    fields: list[TemplateField] = []
    for col in range(1, 500):
        raw = ws.cell(row=header_row_index, column=col).value
        if raw is None:
            continue
        name = str(raw).strip()
        if not name:
            continue
        required = name.lstrip().startswith("*")
        key = slugify(name, separator="_")
        if not key:
            key = f"col_{col}"
        fields.append(TemplateField(field_key=key, field_name=name, column_index=col, required=required))
    return fields


def to_fields_json(fields: list[TemplateField]) -> list[dict[str, Any]]:
    return [
        {
            "field_key": f.field_key,
            "field_name": f.field_name,
            "column_index": f.column_index,
            "required": bool(f.required),
        }
        for f in fields
    ]


def generate_fields_markdown(*, fields: list[TemplateField], out_path: str) -> None:
    p = Path(out_path)
    p.parent.mkdir(parents=True, exist_ok=True)
    lines: list[str] = []
    lines.append("# 导出模板字段（解析自 Excel）")
    lines.append("")
    lines.append("| 列号 | field_key | field_name | 必填 |")
    lines.append("|---:|---|---|:---:|")
    for f in fields:
        lines.append(f"| {f.column_index} | `{f.field_key}` | {f.field_name} | {'✅' if f.required else ''} |")
    lines.append("")
    p.write_text("\n".join(lines), encoding="utf-8")

