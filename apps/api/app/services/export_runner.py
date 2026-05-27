from __future__ import annotations

import re
import uuid
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.config import get_settings
from app.models.default_rule import DefaultRule
from app.models.export_batch import ExportBatch
from app.models.export_field_draft import ExportFieldDraft
from app.models.export_field_mapping import ExportFieldMapping
from app.models.export_record import ExportRecord
from app.models.export_template import ExportTemplate
from app.models.product_asset import ProductAsset
from app.models.product_ai_result import ProductAIResult
from app.models.product_task import ProductTask
from app.services.export_fields import apply_default_rules_with_options
from app.services.export_adapters.registry import get_export_adapter
from app.services.listing_validation import parse_sku_rows_from_draft, validate_task_payload
from app.services.temu_upload_template import normalize_listing_fields
from app.services.export_templates import get_default_export_template


_HTTP_RE = re.compile(r"^https?://", re.IGNORECASE)


@dataclass(frozen=True)
class PreviewRow:
    task_id: int
    export_fields: dict[str, Any]
    field_sources: dict[str, Any]
    validation: dict[str, Any]


@dataclass(frozen=True)
class AssetInfo:
    asset_id: int
    slot: str
    public_url: str
    width: int | None
    height: int | None


def preview_exports(
    session: Session,
    *,
    product_task_ids: list[int],
    template_id: int | None = None,
    default_rule_id: int | None = None,
    adapter_key: str | None = None,
) -> dict[str, Any]:
    adapter = get_export_adapter(adapter_key)
    template_meta = adapter.get_template_meta()

    rows: list[PreviewRow] = []
    for task_id in product_task_ids:
        task = session.get(ProductTask, task_id)
        if task is None:
            continue
        draft = session.scalar(select(ExportFieldDraft).where(ExportFieldDraft.product_task_id == task_id).limit(1))
        if draft is None or default_rule_id is not None:
            draft = apply_default_rules_with_options(session, task=task, selected_rule_id=default_rule_id)
        export_fields, field_sources, validation = _build_row_v2(session, task=task, draft=draft, template_meta=template_meta)
        rows.append(PreviewRow(task_id=task_id, export_fields=export_fields, field_sources=field_sources, validation=validation))

    return {
        "template": {
            "id": None,
            "name": adapter.display_name,
            "version": adapter.adapter_key,
            "platform": "temu",
            "header_row_index": template_meta.get("header_row"),
            "file_path": template_meta.get("template_file_path"),
            "adapter_key": adapter.adapter_key,
        },
        "rows": [
            {
                "product_task_id": r.task_id,
                "export_fields": r.export_fields,
                "field_sources": r.field_sources,
                "validation_result": r.validation,
            }
            for r in rows
        ],
    }


def run_exports(session: Session, *, product_task_ids: list[int]) -> dict[str, Any]:
    return run_exports_with_options(
        session,
        product_task_ids=product_task_ids,
        template_id=None,
        default_rule_id=None,
        adapter_key="miaoshou_temu_non_apparel",
        export_only_valid=True,
    )


def run_exports_with_options(
    session: Session,
    *,
    product_task_ids: list[int],
    template_id: int | None,
    default_rule_id: int | None,
    adapter_key: str | None,
    export_only_valid: bool,
) -> dict[str, Any]:
    adapter = get_export_adapter(adapter_key)
    rule = session.get(DefaultRule, default_rule_id) if default_rule_id else None
    template_meta = adapter.get_template_meta()

    batch_no = _new_batch_no()
    batch = ExportBatch(
        batch_no=batch_no,
        template_id=template_id,
        template_version=adapter.adapter_key,
        export_mode="default_rule",
        default_rule_id=rule.id if rule else None,
        default_rule_name=rule.name if rule else None,
        total_count=len(product_task_ids),
        sku_row_count=len(product_task_ids),
        success_count=0,
        failed_count=0,
        status="running",
    )
    session.add(batch)
    session.commit()
    session.refresh(batch)

    # Build all rows first and persist export_records
    built_rows: list[dict[str, Any]] = []
    blocking_errors: list[str] = []
    common_fields: dict[str, Any] = {}
    for task_id in product_task_ids:
        task = session.get(ProductTask, task_id)
        if task is None:
            continue
        draft = session.scalar(select(ExportFieldDraft).where(ExportFieldDraft.product_task_id == task_id).limit(1))
        if draft is None or default_rule_id is not None:
            draft = apply_default_rules_with_options(session, task=task, selected_rule_id=default_rule_id)
        export_fields, field_sources, validation = _build_row_v2(session, task=task, draft=draft, template_meta=template_meta)
        if not common_fields:
            common_fields = {field: export_fields.get(field, "") for field in template_meta.get("common_field_names", [])}
        status = "success" if not (validation.get("errors") or []) else "failed"
        if status == "failed":
            blocking_errors.append(f"task_id={task_id} 缺少必填字段或字段校验失败")
        rec = ExportRecord(
            product_task_id=task_id,
            template_id=template_id,
            template_version=adapter.adapter_key,
            export_fields_json=export_fields,
            field_sources_json=field_sources,
            validation_result_json=validation,
            status=status,
            batch_no=batch_no,
        )
        session.add(rec)
        session.commit()
        session.refresh(rec)
        if status == "success":
            built_rows.append(export_fields)
        if status == "success":
            batch.success_count += 1
        else:
            batch.failed_count += 1
        session.add(batch)
        session.commit()

    if blocking_errors and not export_only_valid:
        batch.status = "failed"
        session.add(batch)
        session.commit()
        raise RuntimeError("导出前校验未通过：" + "；".join(blocking_errors[:5]))

    if not built_rows:
        batch.status = "failed"
        session.add(batch)
        session.commit()
        raise RuntimeError("没有可导出的通过任务")

    # Generate excel
    exported_rel_path = adapter.write_excel(common_fields=common_fields, rows=built_rows, batch_no=batch_no)
    batch.exported_file_path = exported_rel_path
    batch.status = "success" if batch.failed_count == 0 else "partial_success"
    session.add(batch)
    session.commit()
    session.refresh(batch)

    # backfill records with file path
    session.query(ExportRecord).filter(ExportRecord.batch_no == batch_no).update({"exported_file_path": exported_rel_path})
    session.commit()

    settings = get_settings()
    download_url = f"{settings.public_base_url}/storage/{exported_rel_path}"
    return {"batch_no": batch_no, "batch": batch, "download_url": download_url}


def _validate_export_configuration(*, template: ExportTemplate, mappings: list[ExportFieldMapping]) -> None:
    enabled_mappings = [item for item in mappings if item.enabled]
    if not enabled_mappings:
        raise RuntimeError(f"导出模板 {template.name} 未配置字段映射规则")
    required_mappings = [item for item in enabled_mappings if item.required]
    if not required_mappings:
        raise RuntimeError(f"导出模板 {template.name} 未配置必填字段")


def _build_row_v2(
    session: Session,
    *,
    task: ProductTask,
    draft: ExportFieldDraft | None,
    template_meta: dict[str, Any],
) -> tuple[dict[str, Any], dict[str, Any], dict[str, Any]]:
    fields = normalize_listing_fields((draft.fields_json or {}) if draft is not None else {}, keep_unknown=True, include_auxiliary=True)
    sources = dict((draft.field_sources_json or {}) if draft is not None else {})
    selected_assets = _load_selected_assets(session, task_id=task.id)
    image_settings = _normalize_export_image_settings(getattr(task, "export_image_settings_json", None))

    export_fields: dict[str, Any] = {}
    for field_name in template_meta.get("common_field_names", []):
        export_fields[field_name] = fields.get(field_name, "")
    for field_name in template_meta.get("detail_headers", []):
        value = fields.get(field_name, "")
        if not value and _is_media_field(field_name):
            value = _fill_media_field_by_name(field_name, selected_assets, image_settings=image_settings)
            if value:
                sources[field_name] = {"source": "asset_fallback"}
        export_fields[field_name] = value

    if str(template_meta.get("adapter_key") or "").strip() == "miaoshou_temu_non_apparel":
        _fill_miaoshou_fallback_fields(task=task, export_fields=export_fields)

    if not str(export_fields.get("商品层级") or "").strip():
        export_fields["商品层级"] = "单SKU商品"
    if not str(export_fields.get("SPU货号") or "").strip():
        export_fields["SPU货号"] = task.task_no or task.platform_sku or task.source_id or ""
    if not str(export_fields.get("SKU货号") or "").strip():
        export_fields["SKU货号"] = task.platform_sku or task.source_id or task.task_no or ""
    if not str(export_fields.get("商品名称") or "").strip():
        export_fields["商品名称"] = task.title or ""
    if not str(export_fields.get("英文名称") or "").strip() and isinstance(task.title_package, dict):
        export_fields["英文名称"] = str(task.title_package.get("title_en") or "")

    common_fields = {field: export_fields.get(field, "") for field in template_meta.get("common_field_names", [])}
    detail_row = {field: export_fields.get(field, "") for field in template_meta.get("detail_headers", [])}
    validation = validate_task_payload(
        common_fields=common_fields,
        rows=[detail_row],
        template_meta=template_meta,
        asset_dimensions_by_url={
            asset.public_url: (asset.width, asset.height)
            for asset in selected_assets.values()
            if asset.public_url
        },
    )
    if len(parse_sku_rows_from_draft(fields)) > 1:
        validation["warnings"].append(
            {
                "field": "SKU信息明细",
                "field_name": "SKU信息明细",
                "field_key": "SKU信息明细",
                "type": "multi_sku_not_expanded",
                "message": "当前导出仍按单条主行写入，请复核多 SKU 商品",
                "blocking": False,
            }
        )
    return export_fields, sources, validation


def _fill_miaoshou_fallback_fields(*, task: ProductTask, export_fields: dict[str, Any]) -> None:
    # Keep minimal safe defaults so users can export with existing drafts while migrating.
    if not str(export_fields.get("* 主编号") or "").strip():
        export_fields["* 主编号"] = task.task_no or task.platform_sku or task.source_id or ""
    if not str(export_fields.get("*产品标题") or "").strip():
        export_fields["*产品标题"] = task.title or ""
    if not str(export_fields.get("*英文标题") or "").strip() and isinstance(task.title_package, dict):
        export_fields["*英文标题"] = str(task.title_package.get("title_en") or "")
    if not str(export_fields.get("* 申报价\n（CNY）") or "").strip() and task.price_usd is not None:
        export_fields["* 申报价\n（CNY）"] = str(task.price_usd)


def _build_row(
    session: Session,
    *,
    task: ProductTask,
    template: ExportTemplate,
    mappings: list[ExportFieldMapping],
    draft: ExportFieldDraft | None,
) -> tuple[dict[str, Any], dict[str, Any], dict[str, Any]]:
    export_fields: dict[str, Any] = {}
    field_sources: dict[str, Any] = {}
    errors: list[dict[str, Any]] = []
    warnings: list[dict[str, Any]] = []
    media_field_values: dict[str, str] = {}

    draft_fields = (draft.fields_json or {}) if draft is not None else {}
    url_validation_cache: dict[str, dict[str, Any]] = {}

    selected_assets = _load_selected_assets(session, task_id=task.id)
    image_settings = _normalize_export_image_settings(getattr(task, "export_image_settings_json", None))

    for m in mappings:
        key = m.field_key
        value = None
        source_meta: dict[str, Any] = {"mapping_id": m.id, "source_type": m.source_type, "source_path": m.source_path}
        asset_output_mode = str((m.transform_rule_json or {}).get("asset_output") or "url").strip().lower()

        if not m.enabled:
            value = ""
            source_meta["source_type"] = "disabled"
        elif m.source_type == "draft":
            path = (m.source_path or "").strip()
            path = path.removeprefix("export_field_drafts.").removeprefix("fields.")
            value = draft_fields.get(path)
            if _is_media_field(m.field_name) and not value:
                value = _fill_media_field_by_name(m.field_name or "", selected_assets, image_settings=image_settings)
                if value:
                    source_meta["source_type"] = "asset_fallback"
                    if _is_carousel_media_field(m.field_name):
                        source_meta["size_chart_insert"] = image_settings
        elif m.source_type == "asset":
            slots = _parse_slots(m.source_path or "")
            if _is_carousel_media_field(m.field_name) and any(slot.startswith("carousel_") for slot in slots):
                infos = _arranged_carousel_assets(selected_assets, image_settings=image_settings)
                source_meta["size_chart_insert"] = image_settings
            else:
                infos = [selected_assets.get(s) for s in slots if selected_assets.get(s)]
            infos = [item for item in infos if item]
            value = _format_asset_output(infos, m.transform_rule_json or {})
        elif m.source_type == "task":
            value = getattr(task, (m.source_path or "").strip(), None)
        elif m.source_type == "fixed":
            value = m.default_value
        else:
            value = None

        if (value is None or value == "") and m.default_value:
            value = m.default_value
            source_meta["source_type"] = source_meta.get("source_type") or "fixed"
            source_meta["default_applied"] = True

        export_fields[key] = value if value is not None else ""
        field_sources[key] = source_meta
        if _is_media_field(m.field_name):
            media_field_values[m.field_name or key] = str(export_fields[key] or "")

        # validations
        if m.required and (value is None or (isinstance(value, str) and not value.strip())):
            errors.append({"field_key": key, "field_name": m.field_name, "type": "missing", "message": "required field missing"})

        if m.source_type == "asset" or _is_media_field(m.field_name):
            if value:
                parts = [p.strip() for p in str(value).split(",") if p.strip()]
                if asset_output_mode != "asset_id":
                    bad = [p for p in parts if not _HTTP_RE.match(p)]
                    if bad:
                        errors.append({"field_key": key, "field_name": m.field_name, "type": "invalid_url", "message": "image url invalid", "bad": bad[:3]})
                    url_errors, url_warnings = _validate_remote_media_urls(parts, url_validation_cache)
                    if url_errors:
                        errors.extend(
                            [{"field_key": key, "field_name": m.field_name, **item} for item in url_errors]
                        )
                    if url_warnings:
                        warnings.extend(
                            [{"field_key": key, "field_name": m.field_name, **item} for item in url_warnings]
                        )
                if "轮播图" in (m.field_name or ""):
                    if len(parts) < 3:
                        errors.append({"field_key": key, "field_name": m.field_name, "type": "carousel_min_count", "message": "carousel images require at least 3 urls"})
                    if len(parts) > 10:
                        warnings.append({"field_key": key, "field_name": m.field_name, "type": "carousel_max_count", "message": "carousel images exceed 10 urls"})
                    dim_warnings = _validate_asset_dimensions(parts, selected_assets)
                    warnings.extend(dim_warnings)
            elif m.required:
                errors.append({"field_key": key, "field_name": m.field_name, "type": "missing_image", "message": "image url missing"})

    lang_errors, lang_warnings = _validate_multilingual_carousel_diff(media_field_values)
    errors.extend(lang_errors)
    warnings.extend(lang_warnings)
    warnings.extend(_validate_task_ai_readiness(session, task=task, export_fields=export_fields))

    return export_fields, field_sources, {"errors": errors, "warnings": warnings}


def _validate_task_ai_readiness(
    session: Session,
    *,
    task: ProductTask,
    export_fields: dict[str, Any],
) -> list[dict[str, Any]]:
    warnings: list[dict[str, Any]] = []
    mode = str(task.generation_mode or "").strip()
    ai = session.scalar(select(ProductAIResult).where(ProductAIResult.task_id == task.id).limit(1))

    title_output = (ai.title_package or {}).get("output") if isinstance(ai, ProductAIResult) and isinstance(ai.title_package, dict) else None
    product_info_output = (ai.product_info or {}).get("output") if isinstance(ai, ProductAIResult) and isinstance(ai.product_info, dict) else None
    category_output = (ai.category_match or {}).get("output") if isinstance(ai, ProductAIResult) and isinstance(ai.category_match, dict) else None

    if mode not in {"task_only", "no_ai"}:
        if not isinstance(title_output, dict) or not (title_output.get("title_cn") or title_output.get("title_en")):
            warnings.append(
                {
                    "field_key": "title_package",
                    "type": "ai_title_missing",
                    "message": "AI 标题包未生成或为空；导出字段可能使用原始标题或人工值兜底",
                }
            )
        if not isinstance(category_output, dict) or not (category_output.get("best_path") or category_output.get("selected_category")):
            warnings.append(
                {
                    "field_key": "category_path",
                    "type": "category_recall_missing",
                    "message": "类目召回结果为空；请确认采用类目是否正确",
                }
            )

    if bool(getattr(task, "include_product_info", False)) and mode not in {"task_only", "no_ai"}:
        if not isinstance(product_info_output, dict) or not product_info_output:
            warnings.append(
                {
                    "field_key": "product_info",
                    "type": "product_info_missing",
                    "message": "已开启商品理解，但商品理解结果为空；标题和四宫格可能缺少商品摘要上下文",
                }
            )

    if not str(export_fields.get("商品名称") or "").strip() and not str(export_fields.get("英文名称") or "").strip():
        warnings.append(
            {
                "field_key": "商品名称",
                "type": "listing_title_missing",
                "message": "中文标题和英文标题都为空，请先生成标题或手动填写",
            }
        )
    return warnings


def _load_selected_assets(session: Session, *, task_id: int) -> dict[str, AssetInfo]:
    rows = session.scalars(
        select(ProductAsset)
        .where(ProductAsset.product_task_id == task_id, ProductAsset.selected_for_export.is_(True))
        .order_by(ProductAsset.created_at.desc(), ProductAsset.id.desc())
    ).all()
    out: dict[str, AssetInfo] = {}
    for a in rows:
        if a.public_url:
            out[a.slot] = AssetInfo(asset_id=a.id, slot=a.slot, public_url=a.public_url, width=a.width, height=a.height)
    return out


def _parse_slots(source_path: str) -> list[str]:
    s = (source_path or "").strip()
    if not s:
        return []
    if s.startswith("assets."):
        # assets.carousel_1.selected.public_url
        parts = s.split(".")
        if len(parts) >= 2:
            return [parts[1]]
    # comma list
    return [p.strip() for p in s.split(",") if p.strip()]


def _is_media_field(field_name: str | None) -> bool:
    name = str(field_name or "")
    return any(token in name for token in ("图", "图片", "视频", "轮播图", "预览图", "详情图文"))


def _is_carousel_media_field(field_name: str | None) -> bool:
    return "轮播图" in str(field_name or "")


def _normalize_export_image_settings(value: Any) -> dict[str, Any]:
    raw = value if isinstance(value, dict) else {}
    enabled = raw.get("insert_size_chart_in_carousel")
    if enabled is None:
        enabled = raw.get("size_chart_in_carousel")
    try:
        position = int(raw.get("size_chart_position") or 3)
    except (TypeError, ValueError):
        position = 3
    return {
        "insert_size_chart_in_carousel": bool(True if enabled is None else enabled),
        "size_chart_position": max(1, min(position, 10)),
    }


def _arranged_carousel_assets(
    selected_assets: dict[str, AssetInfo],
    *,
    image_settings: dict[str, Any],
) -> list[AssetInfo]:
    carousel_infos = [
        selected_assets[slot]
        for index in range(1, 11)
        for slot in [f"carousel_{index}"]
        if selected_assets.get(slot) and selected_assets[slot].public_url
    ]
    size_info = selected_assets.get("size_chart")
    if not image_settings.get("insert_size_chart_in_carousel") or not size_info or not size_info.public_url:
        return carousel_infos
    insert_index = max(0, min(int(image_settings.get("size_chart_position") or 3) - 1, len(carousel_infos)))
    return [*carousel_infos[:insert_index], size_info, *carousel_infos[insert_index:]]


def _fill_media_field_by_name(
    field_name: str,
    selected_assets: dict[str, AssetInfo],
    *,
    image_settings: dict[str, Any],
) -> str:
    name = field_name.strip()
    if not name:
        return ""
    if "轮播图" in name:
        arranged = _arranged_carousel_assets(selected_assets, image_settings=image_settings)
        # 支持“商品轮播图1..10”按列回填；无序号时回填逗号拼接
        match = re.search(r"轮播图(\d+)", name)
        if match:
            index = int(match.group(1))
            if 1 <= index <= len(arranged):
                return arranged[index - 1].public_url
            return ""
        urls = [info.public_url for info in arranged]
        return ",".join(urls)
    if "预览图" in name:
        for slot in ("preview_1", "sku_1", "carousel_1"):
            info = selected_assets.get(slot)
            if info and info.public_url:
                return info.public_url
    if "详情图文" in name:
        for slot in ("detail_1", "detail", "carousel_1"):
            info = selected_assets.get(slot)
            if info and info.public_url:
                return info.public_url
    if "主图视频" in name or "详情视频" in name:
        for slot in ("video_1", "video"):
            info = selected_assets.get(slot)
            if info and info.public_url:
                return info.public_url
    return ""


def _format_asset_output(infos: list[AssetInfo], transform_rule_json: dict[str, Any]) -> str:
    mode = str((transform_rule_json or {}).get("asset_output") or "url").strip().lower()
    if mode == "asset_id":
        return ",".join([str(item.asset_id) for item in infos if item.asset_id])
    return ",".join([item.public_url for item in infos if item.public_url])


def _validate_asset_dimensions(urls: list[str], selected_assets: dict[str, AssetInfo]) -> list[dict[str, Any]]:
    url_to_asset = {asset.public_url: asset for asset in selected_assets.values()}
    warnings: list[dict[str, Any]] = []
    for url in urls:
        info = url_to_asset.get(url)
        if info is None:
            continue
        if info.width is not None and info.width < 800:
            warnings.append({"type": "image_width", "message": "image width < 800px", "url": url})
        if info.height is not None and info.height < 800:
            warnings.append({"type": "image_height", "message": "image height < 800px", "url": url})
        if info.width and info.height and info.width != info.height:
            warnings.append({"type": "image_ratio", "message": "image ratio is not 1:1", "url": url})
    return warnings


def _validate_remote_media_urls(
    urls: list[str],
    cache: dict[str, dict[str, Any]],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    errors: list[dict[str, Any]] = []
    warnings: list[dict[str, Any]] = []
    for url in urls:
        if url in cache:
            result = cache[url]
        else:
            result = _probe_remote_url(url)
            cache[url] = result

        if not result.get("ok"):
            errors.append(
                {
                    "type": "url_unreachable",
                    "message": "image url unreachable",
                    "url": url,
                    "detail": result.get("detail"),
                }
            )
            continue

        content_length = result.get("content_length")
        if isinstance(content_length, int) and content_length > 2 * 1024 * 1024:
            errors.append(
                {
                    "type": "image_too_large",
                    "message": "image file size exceeds 2MB",
                    "url": url,
                    "content_length": content_length,
                }
            )
        if content_length is None:
            warnings.append(
                {
                    "type": "content_length_unknown",
                    "message": "image content-length unavailable",
                    "url": url,
                }
            )
    return errors, warnings


def _validate_multilingual_carousel_diff(
    media_field_values: dict[str, str],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    errors: list[dict[str, Any]] = []
    warnings: list[dict[str, Any]] = []
    english_values: list[str] = []
    non_english_values: dict[str, list[str]] = {}

    for field_name, raw_value in media_field_values.items():
        if "轮播图" not in field_name:
            continue
        values = [item.strip() for item in str(raw_value).split(",") if item.strip()]
        if not values:
            continue
        lang = _extract_language_tag(field_name)
        if lang in ("英语", "english", "en", ""):
            english_values.extend(values)
        else:
            non_english_values.setdefault(lang, []).extend(values)

    if not english_values or not non_english_values:
        return errors, warnings

    english_set = set(english_values)
    for lang, items in non_english_values.items():
        diff = [item for item in items if item not in english_set]
        if not diff:
            warnings.append(
                {
                    "type": "multilingual_same_images",
                    "message": f"{lang} 轮播图与英语轮播图完全相同，建议至少 1 张不同",
                }
            )
    return errors, warnings


def _extract_language_tag(field_name: str) -> str:
    if "-" not in field_name:
        return ""
    return field_name.split("-")[-1].strip().lower()


def _probe_remote_url(url: str) -> dict[str, Any]:
    req = Request(
        url,
        method="HEAD",
        headers={"User-Agent": "ai-caiji-export-validator/1.0"},
    )
    try:
        with urlopen(req, timeout=5) as resp:  # noqa: S310
            content_length = _parse_content_length(resp.headers.get("Content-Length"))
            status = getattr(resp, "status", 200)
            if status >= 400:
                return {"ok": False, "detail": f"HTTP {status}"}
            return {"ok": True, "content_length": content_length}
    except HTTPError as exc:
        if exc.code in (405, 501):
            # Some hosts do not allow HEAD; fallback to minimal GET.
            return _probe_remote_url_with_get(url)
        return {"ok": False, "detail": f"HTTP {exc.code}"}
    except URLError as exc:
        return {"ok": False, "detail": str(exc.reason)}
    except Exception as exc:  # noqa: BLE001
        return {"ok": False, "detail": str(exc)}


def _probe_remote_url_with_get(url: str) -> dict[str, Any]:
    req = Request(
        url,
        method="GET",
        headers={
            "User-Agent": "ai-caiji-export-validator/1.0",
            "Range": "bytes=0-0",
        },
    )
    try:
        with urlopen(req, timeout=5) as resp:  # noqa: S310
            content_length = _parse_content_length(resp.headers.get("Content-Length"))
            status = getattr(resp, "status", 200)
            if status >= 400:
                return {"ok": False, "detail": f"HTTP {status}"}
            return {"ok": True, "content_length": content_length}
    except Exception as exc:  # noqa: BLE001
        return {"ok": False, "detail": str(exc)}


def _parse_content_length(value: str | None) -> int | None:
    if not value:
        return None
    try:
        n = int(value)
    except Exception:  # noqa: BLE001
        return None
    return n if n >= 0 else None


def _write_excel(
    *,
    template: ExportTemplate,
    mappings: list[ExportFieldMapping],
    rows: list[tuple[int, dict[str, Any], dict[str, Any], dict[str, Any]]],
    batch_no: str,
) -> str:
    wb = load_workbook(template.file_path)
    ws = wb["Sheet1"] if "Sheet1" in wb.sheetnames else wb[wb.sheetnames[0]]
    start_row = int(template.header_row_index) + 1

    col_by_key = {m.field_key: m.column_index for m in mappings}

    for idx, (_task_id, export_fields, _sources, _validation) in enumerate(rows):
        r = start_row + idx
        for key, value in export_fields.items():
            col = col_by_key.get(key)
            if not col:
                continue
            ws.cell(row=r, column=col).value = value

    settings = get_settings()
    root = Path(settings.storage_root)
    out_dir = root / settings.exports_dir_name
    out_dir.mkdir(parents=True, exist_ok=True)
    filename = f"temu-export-{batch_no}.xlsx"
    out_path = out_dir / filename
    wb.save(out_path)
    # return storage relative path for static mount
    rel = f"{settings.exports_dir_name}/{filename}"
    return rel


def _new_batch_no() -> str:
    return datetime.now(timezone.utc).strftime("%Y%m%d%H%M%S") + "-" + uuid.uuid4().hex[:8]
