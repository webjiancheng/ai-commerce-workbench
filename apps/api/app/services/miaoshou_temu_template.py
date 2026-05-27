from __future__ import annotations

from functools import lru_cache
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


TEMPLATE_SHEET_NAME = "Sheet1"
TEMPLATE_HEADER_ROW = 1
TEMPLATE_HINT_ROW = 2
TEMPLATE_DATA_START_ROW = 3


def get_builtin_miaoshou_temu_template_path() -> str:
    return str(Path(__file__).resolve().parents[4] / "妙手Temu导入模板-非服饰类模板 .xlsx")


@lru_cache(maxsize=1)
def get_miaoshou_temu_template_meta() -> dict[str, Any]:
    return parse_miaoshou_temu_template_meta(template_file_path=get_builtin_miaoshou_temu_template_path())


def parse_miaoshou_temu_template_meta(*, template_file_path: str) -> dict[str, Any]:
    wb = load_workbook(template_file_path, data_only=True)
    ws = wb[TEMPLATE_SHEET_NAME] if TEMPLATE_SHEET_NAME in wb.sheetnames else wb[wb.sheetnames[0]]

    detail_fields: list[dict[str, Any]] = []
    required_fields: list[str] = []

    for col_idx in range(1, ws.max_column + 1):
        header_name = str(ws.cell(row=TEMPLATE_HEADER_ROW, column=col_idx).value or "").strip()
        if not header_name:
            continue
        hint = str(ws.cell(row=TEMPLATE_HINT_ROW, column=col_idx).value or "").strip()
        required = header_name.lstrip().startswith("*")
        if required:
            required_fields.append(header_name)
        detail_fields.append(
            {
                "field_name": header_name,
                "column_index": col_idx,
                "required": required,
                "hint": hint,
            }
        )

    detail_headers = [item["field_name"] for item in detail_fields]
    return {
        "adapter_key": "miaoshou_temu_non_apparel",
        "template_file_path": template_file_path,
        "sheet_name": ws.title,
        "header_row": TEMPLATE_HEADER_ROW,
        "hint_row": TEMPLATE_HINT_ROW,
        "data_start_row": TEMPLATE_DATA_START_ROW,
        "common_fields": [],
        "common_field_names": [],
        "detail_fields": detail_fields,
        "detail_headers": detail_headers,
        "required_fields": sorted(set(required_fields)),
        "field_hints": {item["field_name"]: item.get("hint", "") for item in detail_fields},
        "all_field_names": detail_headers,
    }
