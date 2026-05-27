from __future__ import annotations

from typing import Any

from openpyxl import load_workbook

from app.core.config import get_settings
from app.services.export_adapters.base import ensure_export_dir
from app.services.miaoshou_temu_template import (
    get_builtin_miaoshou_temu_template_path,
    parse_miaoshou_temu_template_meta,
)


class MiaoshouTemuNonApparelAdapter:
    adapter_key = "miaoshou_temu_non_apparel"
    display_name = "妙手Temu非服饰导入模板"

    def get_template_meta(self) -> dict[str, Any]:
        return parse_miaoshou_temu_template_meta(
            template_file_path=get_builtin_miaoshou_temu_template_path()
        )

    def write_excel(
        self,
        *,
        common_fields: dict[str, Any],
        rows: list[dict[str, Any]],
        batch_no: str,
        template_file_path: str | None = None,
        preferred_sheet_name: str | None = None,
    ) -> str:
        template_path = template_file_path or get_builtin_miaoshou_temu_template_path()
        meta = parse_miaoshou_temu_template_meta(template_file_path=template_path)

        wb = load_workbook(template_path)
        sheet_name = preferred_sheet_name if preferred_sheet_name and preferred_sheet_name in wb.sheetnames else meta["sheet_name"]
        ws = wb[sheet_name]

        detail_field_map = {item["field_name"]: item["column_index"] for item in meta["detail_fields"]}
        start_row = int(meta["data_start_row"])
        for row_index, row in enumerate(rows, start=0):
            excel_row = start_row + row_index
            for field_name, value in row.items():
                column_index = detail_field_map.get(field_name)
                if not column_index:
                    continue
                ws.cell(row=excel_row, column=column_index).value = value

        settings = get_settings()
        out_dir = ensure_export_dir(settings.storage_root, settings.exports_dir_name)
        filename = f"{self.adapter_key}-{batch_no}.xlsx"
        out_path = out_dir / filename
        wb.save(out_path)
        return f"{settings.exports_dir_name}/{filename}"
